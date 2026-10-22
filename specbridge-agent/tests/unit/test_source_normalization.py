from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.services import source_normalization


class FakeOcr:
    def __init__(self, text: str = "Whiteboard: login -> OTP") -> None:
        self.text = text
        self.calls: list[str] = []

    async def transcribe(self, data: bytes, *, mime_type: str) -> str:
        self.calls.append(mime_type)
        return self.text


async def normalize(filename: str, data: bytes, ocr=None):
    return await source_normalization.normalize_uploaded_source(
        filename=filename,
        data=data,
        content_type=None,
        author="po-1",
        max_pdf_pages=10,
        max_docx_uncompressed_bytes=10 * 1024 * 1024,
        ocr=ocr,
        max_ocr_pages=5,
    )


@pytest.mark.asyncio
async def test_normalizes_utf8_text() -> None:
    source = await normalize("notes.md", b"Judul\nKebutuhan pengguna.")

    assert source.source_type == "manual_note"
    assert source.title == "notes.md"
    assert source.content == "Judul\nKebutuhan pengguna."
    assert source.metadata["normalized_in"] == "specbridge-agent"


@pytest.mark.asyncio
async def test_normalizes_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Pengguna dapat masuk.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Prioritas"
    table.cell(0, 1).text = "Must"
    stream = BytesIO()
    document.save(stream)

    source = await normalize("requirements.docx", stream.getvalue())

    assert source.source_type == "docx"
    assert "[Paragraph 1]\nPengguna dapat masuk." in source.content
    assert "[Table 1]\nPrioritas | Must" in source.content
    assert source.metadata["paragraph_count"] == 1
    assert source.metadata["locator_markers"] == "paragraph"


@pytest.mark.asyncio
async def test_normalizes_pdf_with_page_markers(monkeypatch) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        is_encrypted = False

        def __init__(self, _stream: BytesIO) -> None:
            self.pages = [Page("Kebutuhan pertama."), Page("Kebutuhan kedua.")]

    monkeypatch.setattr(source_normalization, "PdfReader", Reader)

    source = await normalize("requirements.pdf", b"pdf bytes")

    assert source.source_type == "pdf"
    assert source.content == (
        "[Page 1]\nKebutuhan pertama.\n\n[Page 2]\nKebutuhan kedua."
    )
    assert source.metadata["page_count"] == 2
    assert source.metadata["locator_markers"] == "page"


@pytest.mark.asyncio
async def test_rejects_scanned_pdf_without_text(monkeypatch) -> None:
    class Page:
        def extract_text(self) -> str:
            return ""

    class Reader:
        is_encrypted = False

        def __init__(self, _stream: BytesIO) -> None:
            self.pages = [Page()]

    monkeypatch.setattr(source_normalization, "PdfReader", Reader)

    with pytest.raises(ValueError, match="require OCR"):
        await normalize("scan.pdf", b"pdf bytes")


@pytest.mark.asyncio
async def test_rejects_unsupported_extension() -> None:
    with pytest.raises(ValueError, match="supported upload types"):
        await normalize("recording.mp3", b"audio")


PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


def _reader_with(monkeypatch, texts: list[str]) -> None:
    class Page:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        is_encrypted = False

        def __init__(self, _stream: BytesIO) -> None:
            self.pages = [Page(text) for text in texts]

    monkeypatch.setattr(source_normalization, "PdfReader", Reader)
    monkeypatch.setattr(
        source_normalization, "_single_page_pdf", lambda _reader, index: b"page"
    )


@pytest.mark.asyncio
async def test_ocr_reads_image_uploads() -> None:
    ocr = FakeOcr()

    source = await normalize("whiteboard.png", PNG, ocr=ocr)

    assert source.source_type == "image"
    assert source.content == "[Image]\nWhiteboard: login -> OTP"
    assert source.metadata["extraction"] == "ocr"
    assert ocr.calls == ["image/png"]


@pytest.mark.asyncio
async def test_image_without_ocr_is_rejected() -> None:
    with pytest.raises(ValueError, match="need OCR"):
        await normalize("screenshot.png", PNG)


@pytest.mark.asyncio
async def test_image_with_wrong_content_is_rejected() -> None:
    with pytest.raises(ValueError, match="not a valid PNG"):
        await normalize("fake.jpg", b"not an image", ocr=FakeOcr())


@pytest.mark.asyncio
async def test_ocr_fills_only_scanned_pdf_pages(monkeypatch) -> None:
    _reader_with(monkeypatch, ["Typed page.", ""])
    ocr = FakeOcr("Scanned page text.")

    source = await normalize("mixed.pdf", b"pdf bytes", ocr=ocr)

    assert source.content == "[Page 1]\nTyped page.\n\n[Page 2]\nScanned page text."
    assert source.metadata["ocr_pages"] == [2]
    assert source.metadata["extraction"] == "text_layer+ocr"
    assert ocr.calls == ["application/pdf"]


@pytest.mark.asyncio
async def test_unreadable_scan_is_rejected(monkeypatch) -> None:
    _reader_with(monkeypatch, [""])

    with pytest.raises(ValueError, match="even after OCR"):
        await normalize("blank.pdf", b"pdf bytes", ocr=FakeOcr("[unreadable]"))


@pytest.mark.asyncio
async def test_ocr_page_limit(monkeypatch) -> None:
    _reader_with(monkeypatch, [""] * 6)

    with pytest.raises(ValueError, match="OCR is limited to 5 pages"):
        await normalize("long-scan.pdf", b"pdf bytes", ocr=FakeOcr())
