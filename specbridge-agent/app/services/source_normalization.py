from __future__ import annotations

import asyncio
import re
import uuid
from io import BytesIO
from pathlib import PurePath
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader, PdfWriter

from app.domain.schemas import NormalizedSource, SourceType
from app.infrastructure.ocr import OcrEngine

SUPPORTED_EXTENSIONS: dict[str, SourceType] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "manual_note",
    ".md": "manual_note",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
}

# Gemini takes at most ~20 MB of inline data per request
MAX_OCR_INPUT_BYTES = 18 * 1024 * 1024


async def normalize_uploaded_source(
    *,
    filename: str,
    data: bytes,
    content_type: str | None,
    author: str | None,
    max_pdf_pages: int,
    max_docx_uncompressed_bytes: int,
    ocr: OcrEngine | None = None,
    max_ocr_pages: int = 30,
) -> NormalizedSource:
    safe_name = PurePath(filename.replace("\\", "/")).name.strip()
    if not safe_name:
        raise ValueError("uploaded file must have a filename")

    extension = PurePath(safe_name).suffix.lower()
    source_type = SUPPORTED_EXTENSIONS.get(extension)
    if source_type is None:
        raise ValueError(
            "supported upload types are .pdf, .docx, .txt, .md, .png, .jpg, .jpeg, "
            "and .webp"
        )

    metadata: dict[str, object] = {
        "original_name": safe_name,
        "original_size_bytes": len(data),
        "content_type": content_type or "application/octet-stream",
        "normalized_in": "specbridge-agent",
    }
    if extension == ".pdf":
        content, page_count, ocr_pages = await _extract_pdf(
            data, max_pdf_pages=max_pdf_pages, ocr=ocr, max_ocr_pages=max_ocr_pages
        )
        metadata["page_count"] = page_count
        metadata["locator_markers"] = "page"
        if ocr_pages:
            metadata["ocr_pages"] = ocr_pages
            metadata["extraction"] = (
                "text_layer+ocr" if len(ocr_pages) < page_count else "ocr"
            )
    elif source_type == "image":
        content = await _extract_image(safe_name, data, ocr=ocr)
        metadata["locator_markers"] = "image"
        metadata["extraction"] = "ocr"
    elif extension == ".docx":
        content, paragraph_count = _extract_docx(
            data,
            max_uncompressed_bytes=max_docx_uncompressed_bytes,
        )
        metadata["paragraph_count"] = paragraph_count
        metadata["locator_markers"] = "paragraph"
    else:
        content = _extract_text(data)

    if not content.strip():
        raise ValueError(f"{safe_name} does not contain extractable text")

    stem = re.sub(r"[^a-z0-9]+", "-", PurePath(safe_name).stem.lower()).strip("-")
    source_id = f"upload-{stem[:80] or 'document'}-{uuid.uuid4().hex[:12]}"
    return NormalizedSource(
        source_id=source_id,
        source_type=source_type,
        title=safe_name,
        author=author,
        content=content,
        metadata=metadata,
    )


def _extract_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig").strip()
    except UnicodeDecodeError as exc:
        raise ValueError("text uploads must use UTF-8 encoding") from exc


async def _extract_pdf(
    data: bytes,
    *,
    max_pdf_pages: int,
    ocr: OcrEngine | None,
    max_ocr_pages: int,
) -> tuple[str, int, list[int]]:
    """Text layer first; pages without one (scans) go through OCR when available."""
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("password-protected PDF files are not supported")
        page_count = len(reader.pages)
        if page_count > max_pdf_pages:
            raise ValueError(f"PDF exceeds the {max_pdf_pages}-page limit")
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("PDF could not be read or is malformed") from exc

    blank = [number for number, text in enumerate(texts, start=1) if not text]
    ocr_pages: list[int] = []
    if blank and ocr is not None:
        if len(blank) > max_ocr_pages:
            raise ValueError(
                f"PDF has {len(blank)} scanned pages; OCR is limited to {max_ocr_pages} pages per file"
            )
        transcribed = await asyncio.gather(
            *(
                _transcribe(
                    ocr,
                    _single_page_pdf(reader, number - 1),
                    "application/pdf",
                    f"page {number}",
                )
                for number in blank
            )
        )
        for number, text in zip(blank, transcribed, strict=True):
            if text:
                texts[number - 1] = text
                ocr_pages.append(number)

    pages = [
        f"[Page {number}]\n{text}" for number, text in enumerate(texts, start=1) if text
    ]
    if not pages:
        if ocr is None:
            raise ValueError(
                "PDF has no extractable text; scanned/image PDFs require OCR, which is disabled"
            )
        raise ValueError("PDF has no readable text, even after OCR")
    return "\n\n".join(pages), page_count, ocr_pages


def _single_page_pdf(reader: PdfReader, index: int) -> bytes:
    writer = PdfWriter()
    writer.add_page(reader.pages[index])
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def _image_mime(data: bytes) -> str | None:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


async def _extract_image(name: str, data: bytes, *, ocr: OcrEngine | None) -> str:
    mime_type = _image_mime(data)
    if mime_type is None:
        raise ValueError(f"{name} is not a valid PNG, JPEG, or WebP image")
    if ocr is None:
        raise ValueError("image uploads need OCR, which is disabled on this server")
    text = await _transcribe(ocr, data, mime_type, name)
    if not text:
        raise ValueError(f"{name} has no readable text")
    return f"[Image]\n{text}"


async def _transcribe(ocr: OcrEngine, data: bytes, mime_type: str, label: str) -> str:
    if len(data) > MAX_OCR_INPUT_BYTES:
        raise ValueError(f"{label} is too large for OCR (limit 18 MB)")
    try:
        text = await ocr.transcribe(data, mime_type=mime_type)
    except Exception as exc:
        raise ValueError(
            f"OCR failed for {label}; try again or upload a text version"
        ) from exc
    # A page the model could not read at all comes back as only [unreadable] markers
    if not text.replace("[unreadable]", "").strip():
        return ""
    return text


def _extract_docx(data: bytes, *, max_uncompressed_bytes: int) -> tuple[str, int]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise ValueError("file is not a valid DOCX document")
            expanded_size = sum(item.file_size for item in archive.infolist())
            if expanded_size > max_uncompressed_bytes:
                raise ValueError("DOCX expanded content exceeds the safety limit")
        document = Document(BytesIO(data))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("DOCX could not be read or is malformed") from exc

    blocks: list[str] = []
    paragraph_number = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraph_number += 1
            blocks.append(f"[Paragraph {paragraph_number}]\n{text}")
    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"[Table {table_number}]\n" + "\n".join(rows))

    return "\n\n".join(blocks), paragraph_number
